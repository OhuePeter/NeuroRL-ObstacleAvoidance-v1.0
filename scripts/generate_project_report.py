"""
==========================================================
Generate Project Report: Methods, Code, and Dynamical
Systems (.docx)

Authors:
Peter Ohue
Gunnar Blohm

Description
-----------
Builds a single downloadable Word document that explains the
project end to end for a collaborator or reviewer: the task
and controller, the code pipeline (training, evaluation,
figure generation), the main results with embedded figures,
and the dynamical-systems framework used for the closed-loop
fixed-point analysis. Prose for the task description and the
dynamical-systems primer is reused verbatim from
paper/manuscript_closed_loop_attractor.tex so the voice and
claims stay identical to the manuscript.

Run the figure build scripts first (scripts/build_fig2_trajectories.py,
scripts/build_fig3_velocity.py, scripts/build_fig4_pca.py,
scripts/build_manuscript_attractor_figure.py) so the figures this
report embeds already exist on disk.

Usage
-----
python scripts/generate_project_report.py
==========================================================
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PAPER = Path("paper")
OUTPUT = PAPER / "project_methods_and_dynamics_report.docx"

document = Document()

base_style = document.styles["Normal"]
base_style.font.name = "Calibri"
base_style.font.size = Pt(11)


def heading(text, level=1):
    document.add_heading(text, level=level)


def paragraph(text, bold=False, italic=False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def bullet(text):
    document.add_paragraph(text, style="List Bullet")


def code_block(text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    p.paragraph_format.left_indent = Inches(0.25)
    return p


def figure(path, caption):
    if not path.exists():
        paragraph(f"[Figure not found on disk: {path}. Run the build script first.]", italic=True)
        return
    document.add_picture(str(path), width=Inches(6.2))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = document.add_paragraph()
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9.5)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ============================================================
# Title
# ============================================================

title = document.add_heading(
    "Adaptive Obstacle Avoidance: Methods, Code, and Dynamical-Systems Account",
    level=0,
)
paragraph("Peter Ohue, Gunnar Blohm", italic=True)
paragraph(
    "This document explains the task and controller, the code pipeline that "
    "produces every figure and table in the manuscript, and the "
    "dynamical-systems framework used in the closed-loop fixed-point "
    "analysis (Figure 7). It is intended as a standalone reference for "
    "collaborators and reviewers who want to run or extend the analysis "
    "without reading the full manuscript source."
)

# ============================================================
# 1. Task and controller
# ============================================================

heading("1. Task and controller", level=1)

paragraph(
    "We implemented the obstacle-avoidance environment using the Gymnasium "
    "interface, a standard open-source framework for defining physics-based "
    "reinforcement learning environments. The agent is a two-dimensional "
    "point mass that receives continuous force commands at each timestep "
    "and evolves under Newtonian dynamics within a bounded rectangular "
    "workspace. We deliberately kept the dynamics minimal (a single point "
    "mass with no rotational degrees of freedom and no joint limits) so "
    "that any representational structure found in the trained network is "
    "attributable to learning rather than to the complexity of the "
    "physical system."
)

paragraph(
    "The main perturbation experiment used a corridor formed by two static "
    "circular obstacles positioned symmetrically about the midline of the "
    "workspace. To probe the controller's capacity for online correction, "
    "we applied transient lateral force impulses mid-movement during "
    "evaluation. The impulse magnitude was held constant for six "
    "consecutive timesteps beginning at timestep 50, giving seven "
    "evaluation conditions: an unperturbed control (P0) and three "
    "increasing leftward perturbations (L1, L2, L3) and three increasing "
    "rightward perturbations (R1, R2, R3)."
)

paragraph(
    "All training and evaluation used Stable-Baselines3, an open-source "
    "Python library built on PyTorch. We used Proximal Policy Optimization "
    "(PPO). The policy and value networks are separate feedforward "
    "multilayer perceptrons (MLPs), each with two hidden layers of 256 "
    "units and rectified linear (ReLU) activations. A feedforward MLP "
    "processes each timestep independently, without recurrence or explicit "
    "memory of previous states. Training ran for 3x10^6 environment "
    "timesteps."
)

# ============================================================
# 2. Code pipeline
# ============================================================

heading("2. Code pipeline", level=1)

paragraph(
    "The repository is organised so that each stage of the pipeline, "
    "training, evaluation, neural-population analysis, and figure "
    "generation, has its own entry point. Collaborators can run any stage "
    "independently provided the previous stage's outputs exist on disk."
)

heading("2.1 Training and evaluation", level=2)
bullet("src/training/trainer.py - PPO training loop (Stable-Baselines3).")
bullet("src/environment/ - the Gymnasium obstacle-avoidance environment and observation builder.")
bullet("src/evaluation/evaluator_v2.py - runs the trained policy across all seven perturbation conditions and saves per-episode trajectory, kinematics, and summary CSVs to experiments/version_1_0/results/evaluation_<condition>/.")

heading("2.2 Neural population analysis", level=2)
bullet("src/neural_analysis/activation_extractor.py - extracts the 256-unit hidden-layer activation vector at every evaluation timestep into latent_dataset.npz (activations, condition, episode, timestep).")
bullet("src/neural_analysis/latent_pca.py, latent_clustering.py, latent_correlation.py, latent_trajectory.py - PCA, hierarchical clustering, correlation, and trajectory analyses of the latent dataset.")
bullet("src/neural_analysis/attractor_analysis.py - closed-loop (policy + physics) fixed-point and vector-field analysis in the agent's physical state space, described in Section 3 below.")

heading("2.3 Figure generation (independent, standalone scripts)", level=2)
paragraph(
    "Every main-text figure script is self-contained: it reads directly "
    "from experiments/version_1_0/results/ and writes both a PNG and an "
    "editable SVG (Arial 10pt, COSYNE convention) without depending on any "
    "other figure script."
)
bullet("scripts/build_fig2_trajectories.py -> paper/fig2_trajectories.png + paper/figures/fig2_trajectories.svg")
bullet("scripts/build_fig3_velocity.py -> paper/fig3_velocity.png + paper/figures/fig3_velocity.svg")
bullet("scripts/build_fig4_pca.py -> paper/fig4_pca.png + paper/figures/fig4_pca.svg")
bullet("scripts/attractor_analysis.py -> experiments/version_1_0/results/neural_analysis/attractor/*.png (source panels for Figure 7)")
bullet("scripts/build_manuscript_attractor_figure.py -> paper/fig7_attractor.png + paper/figures/fig7_attractor.svg")
paragraph(
    "Figures 1, 5, and 6 are produced by earlier hand-finished or "
    "script-generated assets already in paper/figures/; Figure 1 in "
    "particular is a hand-drawn biological schematic with no "
    "data-generating script."
)

code_block(
    "python scripts/build_fig2_trajectories.py\n"
    "python scripts/build_fig3_velocity.py\n"
    "python scripts/build_fig4_pca.py\n"
    "python scripts/attractor_analysis.py\n"
    "python scripts/build_manuscript_attractor_figure.py"
)

# ============================================================
# 3. Dynamical-systems framework for the closed loop
# ============================================================

heading("3. Dynamical-systems framework for the closed loop", level=1)

paragraph(
    "PCA, clustering, and decoding describe where the population activity "
    "sits. They do not say why it sits there. To answer that, we treat the "
    "trained policy and the physics engine together as a single "
    "closed-loop dynamical system and analyse it with the standard "
    "vocabulary of dynamical-systems theory."
)

heading("State, trajectory, vector field", level=2)
paragraph(
    "A dynamical system is a rule for how a state changes over time. Here "
    "the state is the agent's physical situation: lateral position x, "
    "forward position y, lateral velocity vx, forward velocity vy. At "
    "every timestep the policy reads the state (plus the fixed goal and "
    "obstacle geometry) and outputs an acceleration; the physics engine "
    "integrates that acceleration into a new velocity and position. "
    "Chaining this over many steps produces a trajectory, exactly the "
    "(x, y) paths plotted in Figure 2. The vector field is the same rule "
    "evaluated everywhere at once: at every point in the state space, it "
    "specifies which direction and how fast the state moves next. A "
    "trajectory is one path through the vector field; the vector field is "
    "the full set of paths the system could take from anywhere."
)

heading("Fixed points, stability, attractors", level=2)
paragraph(
    "A fixed point is a state where the vector field is zero: nothing "
    "changes if the system starts exactly there. In this task, that means "
    "zero lateral velocity and zero commanded lateral acceleration; the "
    "controller has stopped correcting sideways and is content to keep "
    "going straight up the corridor. Whether a fixed point matters "
    "behaviourally depends on its stability. Perturbing the state slightly "
    "and linearising the vector field around the fixed point yields a "
    "matrix, the Jacobian, whose eigenvalues describe what happens next. "
    "In discrete time (the simulation advances in fixed steps of "
    "dt = 0.05 s), an eigenvalue with magnitude below 1 means that "
    "direction decays back to the fixed point, an attractor along that "
    "direction; magnitude above 1 means it grows away, a repeller; a "
    "nonzero imaginary part means the decay oscillates on the way back in "
    "rather than returning in a straight line. A fixed point's basin of "
    "attraction is the set of starting states that eventually settle "
    "there, and two different attractors can coexist in the same system "
    "(multistability) provided their basins do not overlap."
)

heading("Why a feedforward policy still has a closed loop", level=2)
paragraph(
    "The classic version of this analysis (Sussillo and Barak, 2013) finds "
    "fixed points of a recurrent network's own hidden state, because a "
    "recurrent unit's activity at time t+1 is a function of its activity "
    "at time t. Our policy is feedforward: the 256-unit hidden layer "
    "analysed in Figure 4 is a static function of the instantaneous "
    "observation, with no memory of its own previous value, so \u201cthe fixed "
    "points of the hidden layer\u201d is not a well-posed question here. What "
    "is well posed, and what we compute instead, is the fixed points of "
    "the closed loop formed by the policy and the environment's physics, "
    "in the agent's physical state space. This is a legitimate substitute "
    "for the recurrent-network case, and is arguably more interpretable "
    "here, since the state space (x, y, vx, vy) already carries physical "
    "units and meaning, unlike an RNN's abstract hidden units."
)

heading("What the closed-loop analysis found", level=2)
paragraph(
    "Every one of the 140 saved evaluation episodes took the same, left, "
    "detour, so we solved for the fixed point of the reduced lateral "
    "subsystem (x, vx) separately for the left route (matching the "
    "sampled data) and, as a prediction, for the right route. Both fixed "
    "points are locally stable, with two real eigenvalues below 1 in "
    "magnitude and no imaginary component, meaning the lateral state "
    "relaxes back to it smoothly rather than oscillating. The left-route "
    "fixed point sits at x* = 3.32 m with a relaxation time constant of "
    "1.42 s; the right-route fixed point sits at x* = 6.70 m with a time "
    "constant of 1.35 s, close to a mirror image of the left about the "
    "corridor midline at x = 5.0 m. The obstacles sit at x = 4.25 m and "
    "x = 5.75 m with 0.5 m radius, so the left-route fixed point leaves "
    "about 0.43 m of clearance from the near obstacle edge; a rightward "
    "perturbation pushes the agent toward that same near obstacle it is "
    "already closest to, which is the geometric explanation for the "
    "direction-dependent robustness asymmetry reported in the manuscript."
)

# ============================================================
# 4. Results figures
# ============================================================

heading("4. Results figures", level=1)

figure(
    PAPER / "fig2_trajectories.png",
    "Figure 2. Cumulative reach trajectories across all seven perturbation conditions.",
)
figure(
    PAPER / "fig3_velocity.png",
    "Figure 3. (I) Mean lateral velocity and (II) mean speed vs. normalised movement phase, by condition.",
)
figure(
    PAPER / "fig4_pca.png",
    "Figure 4. (I) PC1-PC2 projection of hidden-layer activity coloured by condition. "
    "(II) The same projection coloured by K-means cluster label (K=4).",
)
figure(
    PAPER / "fig7_attractor.png",
    "Figure 7. (I) Closed-loop lateral phase portrait and fixed points per detour route. "
    "(II) Fixed-point location across corridor height vs. empirical detour path. "
    "(III) Fixed-point curves projected into the PC1-PC2 manifold.",
)

# ============================================================
# 5. Reproducibility notes
# ============================================================

heading("5. Reproducibility notes", level=1)
bullet("Run commands as `python -m pytest ...` or `python scripts/...` from the repository root with the project .venv active.")
bullet("Figures 2, 3, 4, and 7 read from experiments/version_1_0/results/ (not version_2_0, which only retains final neural-analysis manuscript PDFs, not raw per-episode recordings).")
bullet(
    "The PCA variance reported for Figure 4 (PC1=68.0%, PC2=25.3%, total=93.3%) is computed "
    "directly from experiments/version_1_0/results/latent_dataset.npz with unit-wise "
    "standardisation before PCA, and is reproducible by running scripts/build_fig4_pca.py. "
    "It differs from the 74.8% figure quoted in the COSYNE abstract and grant proposal, which "
    "was computed by an earlier version_2_0 pipeline whose raw per-episode recordings are no "
    "longer present in this repository; that number is kept as-is in already-submitted "
    "documents and is not retroactively changed here."
)

document.save(OUTPUT)

print(f"Report saved to:\n{OUTPUT.resolve()}")
